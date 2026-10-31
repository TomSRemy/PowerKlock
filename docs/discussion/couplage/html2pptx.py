#!/usr/bin/env python3
"""PowerPoint version of a deck built as HTML (deck_cas.py, deck.py): every slide is measured in Chromium (position,
size, font, colour, bold of each block) and rebuilt in python-pptx at the same place, with editable text and native
tables; drawings and charts are inserted as images. Usage: html2pptx.py /tmp/coupl/deck_cas.html couplage_cas.pptx"""
import sys, base64, io, re
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
PX = 9525                                         # EMU per CSS pixel (96 dpi)

JS = r"""
() => {
  const runs = (el) => {
    // paragraphs of runs: [{text, b, color, size}], one paragraph per block child / <br>
    const paras = [[]];
    const walk = (n) => {
      if (n.nodeType === 3) {
        const t = n.textContent.replace(/\s+/g, ' ');
        if (!t.trim() && !paras[paras.length - 1].length) return;
        const s = getComputedStyle(n.parentElement);
        paras[paras.length - 1].push({text: s.textTransform === 'uppercase' ? t.toUpperCase() : t, b: parseInt(s.fontWeight) >= 600, color: s.color, size: parseFloat(s.fontSize)});
      } else if (n.nodeType === 1) {
        if (n.tagName === 'BR') { paras.push([]); return; }
        const block = ['P', 'LI', 'DIV'].includes(n.tagName);
        if (block && paras[paras.length - 1].length) paras.push([]);
        n.childNodes.forEach(walk);
        if (block) paras.push([]);
      }
    };
    el.childNodes.forEach(walk);
    return paras.filter(p => p.length && p.map(r => r.text).join('').trim());
  };
  const out = [];
  document.querySelectorAll('section').forEach(sec => {
    const b0 = sec.getBoundingClientRect(); const blocks = []; const taken = [];
    const sel = '.k, h1, img, .verdict, p, ul, ol, table, .box, .t, .f';
    sec.querySelectorAll(sel).forEach(el => {
      if (taken.some(t => t.contains(el))) return;
      taken.push(el);
      const r = el.getBoundingClientRect(); const s = getComputedStyle(el);
      const o = {tag: el.tagName, cls: el.className, x: r.left - b0.left, y: r.top - b0.top, w: r.width, h: r.height,
                 size: parseFloat(s.fontSize), color: s.color, bg: s.backgroundColor, align: s.textAlign};
      if (el.tagName === 'IMG') { o.src = el.src; }
      else if (el.tagName === 'TABLE') {
        o.rows = [...el.rows].map(tr => [...tr.cells].map(td => {
          const q = td.getBoundingClientRect(); const cs = getComputedStyle(td);
          return {w: q.width, h: q.height, x: q.left - b0.left, bg: cs.backgroundColor, align: cs.textAlign, paras: runs(td),
                  pl: parseFloat(cs.paddingLeft), pt: parseFloat(cs.paddingTop), span: td.colSpan};
        }));
      } else if (el.tagName === 'UL' || el.tagName === 'OL') {
        o.items = [...el.children].map(li => runs(li)); o.ordered = el.tagName === 'OL';
      } else { o.paras = runs(el); o.pl = parseFloat(s.paddingLeft); o.pt = parseFloat(s.paddingTop); }
      blocks.push(o);
    });
    out.push({bg: getComputedStyle(sec).backgroundColor, cover: sec.classList.contains('cover'), blocks});
  });
  return out;
}
"""

def rgb(c):
    m = re.findall(r'[\d.]+', c or '')
    if len(m) < 3 or (len(m) == 4 and float(m[3]) == 0):
        return None
    return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))

def fill_runs(tf, paras, bullet=None, align=None, size_scale=1.0):
    tf.word_wrap = True
    first = True
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        if align == 'right': p.alignment = PP_ALIGN.RIGHT
        elif align == 'center': p.alignment = PP_ALIGN.CENTER
        if bullet:
            pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(int(16 * PX))); pPr.set('indent', str(int(-14 * PX)))
            if bullet == 'num':
                bu = etree.SubElement(pPr, qn('a:buAutoNum')); bu.set('type', 'arabicPeriod')
            else:
                bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', '•')
            p.space_after = Pt(5)
        text_runs = para
        # trim leading/trailing spaces of the paragraph
        if text_runs:
            text_runs[0] = dict(text_runs[0], text=text_runs[0]['text'].lstrip())
            text_runs[-1] = dict(text_runs[-1], text=text_runs[-1]['text'].rstrip())
        for r in text_runs:
            if not r['text']:
                continue
            run = p.add_run(); run.text = r['text']
            f = run.font; f.size = Pt(r['size'] * 0.75 * size_scale); f.bold = r['b']; f.name = 'Arial'
            c = rgb(r['color'])
            if c is not None: f.color.rgb = c

def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(max(w, 4) * PX)), Emu(int(max(h, 4) * PX)))
    tf = tb.text_frame
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tb, tf

with sync_playwright() as p:
    b = p.chromium.launch(); pg = b.new_page(viewport={'width': 1280, 'height': 720})
    pg.goto('file://' + SRC); pg.wait_for_timeout(500)
    deck = pg.evaluate(JS); b.close()

prs = Presentation(); prs.slide_width = Emu(1280 * PX); prs.slide_height = Emu(720 * PX)
blank = prs.slide_layouts[6]
for sd in deck:
    s = prs.slides.add_slide(blank)
    bgc = rgb(sd['bg'])
    if sd['cover']:
        s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0xF1, 0xF0, 0xFA)
    elif bgc is not None:
        s.background.fill.solid(); s.background.fill.fore_color.rgb = bgc
    for o in sd['blocks']:
        if o['w'] < 1 or o['h'] < 1:
            continue
        if o['tag'] == 'IMG':
            data = base64.b64decode(o['src'].split(',', 1)[1])
            s.shapes.add_picture(io.BytesIO(data), Emu(int(o['x'] * PX)), Emu(int(o['y'] * PX)), Emu(int(o['w'] * PX)), Emu(int(o['h'] * PX)))
        elif o['tag'] == 'TABLE':
            rows = o['rows']; nr = len(rows); nc = max(sum(c['span'] for c in r) for r in rows)
            widths = None
            for r in rows:
                if len(r) == nc:
                    widths = [c['w'] for c in r]; break
            shape = s.shapes.add_table(nr, nc, Emu(int(o['x'] * PX)), Emu(int(o['y'] * PX)), Emu(int(o['w'] * PX)), Emu(int(o['h'] * PX)))
            t = shape.table
            tblPr = t._tbl.tblPr; tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
            for j, wv in enumerate(widths or [o['w'] / nc] * nc):
                t.columns[j].width = Emu(int(wv * PX))
            for i, r in enumerate(rows):
                t.rows[i].height = Emu(int(r[0]['h'] * PX))
                j = 0
                for c in r:
                    cell = t.cell(i, j)
                    if c['span'] > 1:
                        cell.merge(t.cell(i, j + c['span'] - 1))
                    cell.margin_left = cell.margin_right = Emu(int(c['pl'] * PX)); cell.margin_top = cell.margin_bottom = Emu(int(c['pt'] * PX))
                    cb = rgb(c['bg'])
                    cell.fill.solid(); cell.fill.fore_color.rgb = cb if cb is not None else RGBColor(255, 255, 255)
                    cell.text_frame.paragraphs[0].text = ''
                    fill_runs(cell.text_frame, c['paras'] or [[{'text': '', 'b': False, 'color': '', 'size': 12}]], align=c['align'])
                    j += c['span']
        elif o['tag'] in ('UL', 'OL'):
            tb, tf = box(s, o['x'], o['y'], o['w'], o['h'] + 6)
            fill_runs(tf, [sum(item, []) for item in o['items']], bullet='num' if o['ordered'] else 'dot')
        else:
            bgc = rgb(o['bg'])
            if o['cls'] == 'f':                       # footer: wider box, right-aligned, so it never wraps
                o['x'] -= 160; o['w'] += 160; o['align'] = 'right'
            tb, tf = box(s, o['x'], o['y'], o['w'], o['h'] + 4)
            if bgc is not None and bgc != RGBColor(255, 255, 255):
                tb.fill.solid(); tb.fill.fore_color.rgb = bgc
                tf.margin_left = tf.margin_right = Emu(int(max(o.get('pl') or 0, 8) * PX)); tf.margin_top = tf.margin_bottom = Emu(int(max(o.get('pt') or 0, 6) * PX))
            fill_runs(tf, o['paras'], align=o['align'] if o['align'] in ('right', 'center') else None)
prs.save(OUT)
print(len(deck), 'slides →', OUT)
