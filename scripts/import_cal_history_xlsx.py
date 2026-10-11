"""
import_cal_history_xlsx.py — one-off import of a user-supplied Excel archive of
French Cal-year power base forward prices (EEX, digitized from Energy-Charts.info
PDF charts) into data/history/fr_market_calchain.json.

This extends FR's Cal-year history well beyond what the SEFE cockpit API serves
live (expired contracts return no data there — see fetch_fr_market.py's
fetch_country_cal_years docstring), from ~2023 back to 2020-12-28.

For years also covered by the API (2027, 2028, 2029 as of writing), the Excel
data is used up to its own last date and API points strictly after that are
appended, so the series stays current via the daily cron without re-running
this script. Cal+1/+2/+3 rolling series are recomputed from the merged years
with the same calendar-year-boundary stitching as fetch_fr_market.py, so a
newly-merged year (2022-2026) automatically extends the rolling curves too.

Usage:
  python scripts/import_cal_history_xlsx.py <path-to-xlsx>
"""
import sys
import json
from pathlib import Path
from collections import defaultdict

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from fetch_fr_market import build_rolling_cal_series, OUT_CALCHAIN  # noqa: E402

SHEET_PREFIX = "Cal-"


def load_excel_years(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    years = {}
    for name in wb.sheetnames:
        if not name.startswith(SHEET_PREFIX):
            continue
        year = int(name[len(SHEET_PREFIX):])
        ws = wb[name]
        by_date = defaultdict(list)
        for r in range(2, ws.max_row + 1):
            d = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if d is None or v is None:
                continue
            by_date[d.date().isoformat()].append(float(v))
        dates = sorted(by_date)
        values = [sum(by_date[d]) / len(by_date[d]) for d in dates]  # average any duplicate date
        years[year] = {"dates": dates, "values": values}
        print(f"  Cal{year}: {len(dates)} points, {dates[0]} -> {dates[-1]}")
    return years


def merge_year(excel_series, api_series, market_name):
    """Excel data as the historical backbone; any API point strictly after the
    Excel's last date is appended, so the series keeps extending day to day."""
    dates, values = list(excel_series["dates"]), list(excel_series["values"])
    cutoff = dates[-1] if dates else None
    if api_series:
        for d, v in zip(api_series.get("dates", []), api_series.get("values", [])):
            if cutoff is None or d > cutoff:
                dates.append(d)
                values.append(v)
    return {"marketName": market_name, "dates": dates, "values": values}


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/import_cal_history_xlsx.py <path-to-xlsx>")
        sys.exit(1)
    xlsx_path = sys.argv[1]

    print(f"Reading {xlsx_path} ...")
    excel_years = load_excel_years(xlsx_path)

    print(f"Loading existing {OUT_CALCHAIN} ...")
    calchain = json.loads(OUT_CALCHAIN.read_text()) if OUT_CALCHAIN.exists() else {}
    fr = calchain.setdefault("FR", {"years": {}, "rolling": {}})
    existing_years = {int(y): s for y, s in fr.get("years", {}).items()}

    merged_years = {}
    for year in sorted(set(excel_years) | set(existing_years)):
        excel_s = excel_years.get(year)
        api_s = existing_years.get(year)
        if excel_s:
            name = (api_s or {}).get("marketName") or f"FR Cal{year} base"
            merged_years[year] = merge_year(excel_s, api_s, name)
            src = "excel+api" if api_s else "excel only"
        else:
            merged_years[year] = api_s
            src = "api only"
        print(f"  Cal{year}: {len(merged_years[year]['dates'])} points ({src})")

    fr["years"] = {str(y): s for y, s in merged_years.items()}

    print("Recomputing rolling Cal+1/+2/+3 from the merged years...")
    rolling = {}
    for offset, label in ((1, "Cal+1"), (2, "Cal+2"), (3, "Cal+3")):
        rolling[label] = build_rolling_cal_series(merged_years, offset)
        print(f"  {label}: {len(rolling[label]['dates'])} points, "
              f"{rolling[label]['dates'][0]} -> {rolling[label]['dates'][-1]}")
    fr["rolling"] = rolling

    OUT_CALCHAIN.parent.mkdir(parents=True, exist_ok=True)
    OUT_CALCHAIN.write_text(json.dumps(calchain, ensure_ascii=False))
    print(f"Wrote {OUT_CALCHAIN}")


if __name__ == "__main__":
    main()
